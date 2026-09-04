from flask import Flask, request, render_template_string
import openpyxl
import time
import uuid

app = Flask(__name__)

# 短暫紀錄資料庫 (儲存於記憶體，重啟伺服器後清空)
history_db = {}

HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="zh-TW">
<head>
    <meta charset="UTF-8">
    <title>農林作物調查估價表 - 視覺化智能校對系統</title>
    <style>
        body { font-family: '微軟正黑體', sans-serif; background-color: #f4f7f6; margin: 0; padding: 20px; color: #333; }
        .header-title { text-align: center; color: #2c3e50; margin-bottom: 20px; }
        
        /* 版面配置：左側主畫面，右側紀錄面板 */
        .layout-container { display: flex; gap: 20px; max-width: 1300px; margin: auto; align-items: flex-start; }
        .main-content { flex: 1; background: #fff; padding: 30px; border-radius: 10px; box-shadow: 0 4px 15px rgba(0,0,0,0.05); }
        
        .upload-box { border: 2px dashed #95a5a6; padding: 30px; text-align: center; border-radius: 10px; margin-bottom: 20px; background: #fafbfc; }
        .success { color: #27ae60; font-weight: bold; background: #e9f7ef; padding: 10px; border-radius: 5px; }
        .error-list { color: #c0392b; background: #fdf2f0; padding: 15px; border-left: 5px solid #e74c3c; border-radius: 4px; line-height: 1.6; }
        .file-info { font-size: 16px; font-weight: bold; color: #34495e; padding: 10px 0; border-bottom: 2px solid #ecf0f1; margin-bottom: 20px; }

        /* 右側瀏覽紀錄面板 */
        .history-panel { width: 320px; background: #fff; border-radius: 10px; box-shadow: 0 4px 15px rgba(0,0,0,0.05); overflow: hidden; position: sticky; top: 20px; }
        .history-header { background: #3498db; color: #fff; padding: 15px; font-size: 16px; font-weight: bold; display: flex; justify-content: space-between; cursor: pointer; }
        .history-content { max-height: 70vh; overflow-y: auto; background: #fdfdfd; }
        .history-item { padding: 12px 15px; border-bottom: 1px solid #eee; cursor: pointer; display: flex; flex-direction: column; transition: background 0.2s; }
        .history-item:hover { background: #f0f8ff; }
        .history-item.active { background: #e1f0fa; border-left: 4px solid #2980b9; }
        .history-time { font-size: 12px; color: #7f8c8d; margin-top: 4px; }
        
        /* 視覺化圖面專屬 CSS */
        .map-container { overflow-x: auto; margin-top: 20px; border-radius: 8px; border: 1px solid #ddd; }
        .excel-map { border-collapse: collapse; width: 100%; background: #fff; }
        .excel-map th { background-color: #ecf0f1; border: 1px solid #bdc3c7; font-size: 12px; padding: 6px; color: #555; position: sticky; top: 0; z-index: 2; }
        .excel-map td { border: 1px solid #ecf0f1; height: 30px; min-width: 50px; text-align: center; position: relative; font-size: 12px; }
        
        .cell-valid { background-color: #eafaf1; color: #27ae60; }
        .cell-empty { background-color: #f9f9f9; color: #ccc; }
        .cell-error { background-color: #fadbd8; color: #c0392b; border: 2px solid #e74c3c !important; font-weight: bold; cursor: help; }
        
        /* Tooltip 特效 */
        .tooltip {
            position: absolute; background-color: rgba(0,0,0,0.85); color: #fff; padding: 8px 12px;
            border-radius: 6px; font-size: 12px; font-weight: normal; bottom: 120%; left: 50%;
            transform: translateX(-50%); white-space: nowrap; opacity: 0; visibility: hidden;
            transition: opacity 0.2s ease; pointer-events: none; z-index: 9999;
        }
        .cell-error:hover .tooltip { opacity: 1; visibility: visible; }
        .cell-error:hover { box-shadow: 2px 2px 10px rgba(231,76,60,0.4); z-index: 10; }
    </style>
    <script>
        function toggleHistory() {
            var content = document.getElementById('historyContent');
            content.style.display = (content.style.display === 'none') ? 'block' : 'none';
        }
    </script>
</head>
<body>
    <h2 class="header-title">📊 農林作物調查估價表 - 智能視覺化校對系統</h2>
    
    <div class="layout-container">
        <!-- 左側主畫面 -->
        <div class="main-content">
            <div class="upload-box">
                <form method="POST" enctype="multipart/form-data" action="/">
                    <input type="file" name="excel_file" accept=".xlsx" required>
                    <button type="submit" style="padding: 8px 20px; font-weight: bold; cursor: pointer;">上傳並開始校對</button>
                </form>
            </div>

            {% if record %}
                <div class="file-info">📄 本次校對檔案：{{ record.filename }}</div>
                
                {% if record.error_list|length == 0 %}
                    <p class="success">✅ 完美！字體大小與循環格線格式皆完全符合規定。(系統已自動忽略 S 欄之後的內容)</p>
                {% else %}
                    <p>💡 系統已為您重繪 Excel 佈局 (僅校對至 S 欄)，將滑鼠移至<span style="color:#c0392b; font-weight:bold;">紅色格子</span>可查看錯誤原因。</p>
                    
                    <div class="map-container">
                        <table class="excel-map">
                            <tr>
                                <th></th>
                                {% for c in range(1, record.max_col + 1) %}
                                    <th>{{ get_col_letter(c) }}</th>
                                {% endfor %}
                            </tr>
                            {% for r in range(1, record.max_row + 1) %}
                            <tr>
                                <th>{{ r }}</th>
                                {% for c in range(1, record.max_col + 1) %}
                                    {% set cell = record.map_data.get((r, c), {'status': 'empty', 'text': '', 'error': ''}) %}
                                    <td class="cell-{{ cell.status }}">
                                        {{ cell.text }}
                                        {% if cell.status == 'error' %}
                                            <div class="tooltip">{{ cell.error }}</div>
                                        {% endif %}
                                    </td>
                                {% endfor %}
                            </tr>
                            {% endfor %}
                        </table>
                    </div>

                    <div class="error-list" style="margin-top: 20px;">
                        <h4>📋 詳細錯誤報告清單：</h4>
                        <ul>
                            {% for err in record.error_list %}
                                <li>{{ err }}</li>
                            {% endfor %}
                        </ul>
                    </div>
                {% endif %}
            {% endif %}
        </div>

        <!-- 右側歷史紀錄面板 -->
        <div class="history-panel">
            <div class="history-header" onclick="toggleHistory()">
                <span>📂 瀏覽紀錄</span>
                <span>▼</span>
            </div>
            <div class="history-content" id="historyContent">
                {% for uid, hist in history.items()|reverse %}
                    <div class="history-item {% if uid == current_id %}active{% endif %}" onclick="window.location.href='/?id={{ uid }}'">
                        <strong>{{ hist.filename }}</strong>
                        <span class="history-time">{{ hist.time }}</span>
                    </div>
                {% else %}
                    <div style="padding: 15px; text-align: center; color: #999; font-size: 13px;">尚無校對紀錄</div>
                {% endfor %}
            </div>
        </div>
    </div>
</body>
</html>
"""

@app.route("/", methods=["GET", "POST"])
def index():
    record_id = request.args.get('id')
    
    if request.method == "POST":
        file = request.files["excel_file"]
        if file and file.filename.endswith('.xlsx'):
            wb = openpyxl.load_workbook(file, data_only=False)
            ws = wb.active
            
            error_list = []
            map_data = {}
            
            # 【關鍵修正】強制限制最大校對欄位為 19 (即 S 欄)
            max_col = 19 
            max_row = min(ws.max_row, 100) # 為了效能，圖面最多畫 100 列
            merged_ranges = ws.merged_cells.ranges

            for row in range(1, max_row + 1):
                # 這裡的迴圈範圍嚴格被鎖死在 1~19，S 欄以後的資料 (T, U, V...) 完全不會被讀取與報錯
                for col in range(1, max_col + 1):
                    cell = ws.cell(row=row, column=col)
                    coord = cell.coordinate
                    
                    # 判斷合併儲存格與「有效底部列」
                    is_phantom = False
                    eff_bottom_row = row
                    
                    for mr in merged_ranges:
                        min_col, min_row, max_mr_col, mr_max_row = mr.bounds
                        # 若合併範圍超出 S 欄，我們也只在 A~S 的範圍內處理
                        if min_row <= row <= mr_max_row and min_col <= col <= max_mr_col:
                            if row == min_row and col == min_col:
                                is_phantom = False
                                eff_bottom_row = mr_max_row # 讓底線判斷以此格為準
                            else:
                                is_phantom = True
                            break

                    if is_phantom:
                        map_data[(row, col)] = {'status': 'empty', 'text': ''}
                        continue
                        
                    val = cell.value
                    if val is None:
                        map_data[(row, col)] = {'status': 'empty', 'text': ''}
                        continue

                    val_str = str(val).strip()
                    font = cell.font
                    border = cell.border
                    font_size = font.size if font else None
                    errors = []

                    # 取得動態對照基準列 (處理大於等於 10 列的循環)
                    ref_row = 10 + ((row - 10) % 17) if row >= 10 else row
                    eff_ref_row = 10 + ((eff_bottom_row - 10) % 17) if eff_bottom_row >= 10 else eff_bottom_row

                    # ============== 【字體大小檢核】 ==============
                    if col == 1 and row <= 9:
                        pass # 條件 2：A1-A9無須檢核
                    elif row <= 2 and col <= 17:
                        if font_size != 15:
                            errors.append(f"標題字體應為 15pt (目前 {font_size}pt)")
                    elif row == 8 and col == 15:
                        if font_size != 10:
                            errors.append(f"字體應為 10pt (目前 {font_size}pt)")
                    elif 21 <= ref_row <= 25 and 3 <= col <= 19:
                        if font_size != 10:
                            errors.append(f"字體應為 10pt (目前 {font_size}pt)")
                    else:
                        pass # 其餘未定義區域預設放行

                    # ============== 【格線檢核】 ==============
                    # 條件 3：A21-S24、A26-S26 無底部格線 (以有效底部列 eff_ref_row 為準)
                    is_no_bottom_border = (21 <= eff_ref_row <= 24) or (eff_ref_row == 26)
                    
                    if is_no_bottom_border:
                        if border and border.bottom and border.bottom.style:
                            errors.append(f"對照第 {eff_ref_row} 列，不應有底部格線")
                    else:
                        if row >= 10: # 一般明細資料列皆須底線
                            if not border or not border.bottom or not border.bottom.style:
                                errors.append(f"對照第 {eff_ref_row} 列，缺少底部格線")

                    # 圖面預覽文字
                    display_text = val_str[:4] + '..' if len(val_str) > 4 else val_str
                    
                    if errors:
                        err_msg = "、".join(errors)
                        map_data[(row, col)] = {'status': 'error', 'text': display_text, 'error': err_msg}
                        error_list.append(f"❌ {coord}: {err_msg}")
                    else:
                        map_data[(row, col)] = {'status': 'valid', 'text': display_text}

            # 存入記憶體資料庫
            uid = str(uuid.uuid4())[:8]
            history_db[uid] = {
                'filename': file.filename,
                'time': time.strftime("%H:%M:%S"),
                'map_data': map_data,
                'error_list': error_list,
                'max_row': max_row,
                'max_col': max_col
            }
            
            return render_template_string(HTML_TEMPLATE, record=history_db[uid], history=history_db, current_id=uid, get_col_letter=openpyxl.utils.get_column_letter)

    # 點擊歷史紀錄切換
    if record_id and record_id in history_db:
        return render_template_string(HTML_TEMPLATE, record=history_db[record_id], history=history_db, current_id=record_id, get_col_letter=openpyxl.utils.get_column_letter)

    # 首頁預設畫面
    return render_template_string(HTML_TEMPLATE, record=None, history=history_db, current_id=None)

if __name__ == "__main__":
    app.run(debug=True, port=5000)